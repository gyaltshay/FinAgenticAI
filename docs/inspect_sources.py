from pathlib import Path
import zipfile, xml.etree.ElementTree as ET, json
from pypdf import PdfReader
import openpyxl

source=Path(r'C:\Users\user\Downloads')
out=Path(r'D:\FinAgenticAI\docs\source-review')
out.mkdir(exist_ok=True)
names=['TTPL-Proposal for Development of an Agentic AI System for Automated Financial Reporting (1).docx','Revised_Methodology_Agentic_AI_Financial_Reporting.docx','Group4_Proposal Report_BE4IT.pdf','August month.xlsx','Aug PPT.pptx','Slides for Business Review - Q2 2026. final .xlsx']
for name in names:
 p=source/name
 lines=[]
 if p.suffix=='.docx':
  with zipfile.ZipFile(p) as z:
   root=ET.fromstring(z.read('word/document.xml'))
   for para in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
    text=''.join(t.text or '' for t in para.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
    if text: lines.append(text)
 elif p.suffix=='.pdf':
  reader=PdfReader(p)
  for i,page in enumerate(reader.pages):lines.append(f'\nPAGE {i+1}\n'+(page.extract_text() or ''))
 elif p.suffix=='.pptx':
  with zipfile.ZipFile(p) as z:
   slides=sorted([n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')],key=lambda n:int(n.split('slide')[-1].split('.')[0]))
   for slide in slides:
    root=ET.fromstring(z.read(slide))
    lines.append('\n'+slide+'\n'+'\n'.join(t.text or '' for t in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t')))
 elif p.suffix=='.xlsx':
  wb=openpyxl.load_workbook(p,data_only=True)
  for ws in wb:
   lines.append(f'\nSHEET {ws.title} | rows={ws.max_row} columns={ws.max_column}\n')
   for row in ws:
    values=[f'{c.coordinate}={c.value}' for c in row if c.value is not None]
    if values:lines.append(' | '.join(values))
 text='\n'.join(lines)
 dest=out/(p.stem+'.txt')
 dest.write_text(text,encoding='utf-8')
 print(json.dumps({'file':name,'extracted':str(dest),'characters':len(text),'lines':len(lines)}))
