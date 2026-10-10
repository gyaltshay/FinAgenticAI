from pathlib import Path
import openpyxl, zipfile, xml.etree.ElementTree as ET
root=Path(r'C:\Users\user\Downloads')
pairs=[('August month.xlsx',['Month Financials','YTD Financials','Monthly PPR','YTD PR','Balance Sheet']),('Slides for Business Review - Q2 2026. final .xlsx',["Q2'2026 Actuals",'YTD June 2026',"Q3'26 Forecast",'FY 2026 Act Forecast'])]
for name,sheets in pairs:
 print('\nFILE',name)
 wb=openpyxl.load_workbook(root/name,data_only=True)
 for sheet in sheets:
  ws=wb[sheet]
  print('\nSHEET',sheet)
  for row in ws.iter_rows(min_row=1,max_row=36,max_col=14):
   vals=[f'{c.coordinate}={c.value}' for c in row if c.value is not None]
   if vals:print(' | '.join(vals))
 with zipfile.ZipFile(root/name) as z:
  links=[n for n in z.namelist() if n.startswith('xl/externalLinks/') and n.endswith('.xml')]
  print('External link parts:',len(links))
 wbform=openpyxl.load_workbook(root/name,data_only=False)
 for sheet in sheets[:2]:
  ws=wbform[sheet]
  print('Number formats and formulas:',sheet)
  for row in ws.iter_rows(min_row=1,max_row=24,max_col=6):
   for c in row:
    if c.data_type=='f' or (c.column==3 and c.value is not None):print(c.coordinate,str(c.value)[:180],c.number_format)
with zipfile.ZipFile(root/'Aug PPT.pptx') as z:
 media=[n for n in z.namelist() if n.startswith('ppt/media/')]
 print('PPT media:',media)
 for n in media:
  if n.endswith('.png') or n.endswith('.jpeg'):
   (Path(r'D:\FinAgenticAI\docs\source-review')/Path(n).name).write_bytes(z.read(n))
